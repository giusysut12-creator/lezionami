import io

import pytest
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from app.extract import (
    ExtractionError, build_transcript_passages, detect_issues, extract_file, from_pasted_text,
    segment_passages,
)
from tests.conftest import FIXTURES


def _image_only_pdf(pages: int = 2) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    image = Image.new("RGB", (400, 300), (200, 200, 230))
    for _ in range(pages):
        pdf.drawImage(ImageReader(image), 50, 400, 400, 300)
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def _text_pdf(lines: list[str], extra_image_page: bool = False) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    y = 780
    for line in lines:
        pdf.drawString(50, y, line)
        y -= 16
    pdf.showPage()
    if extra_image_page:
        pdf.drawImage(ImageReader(Image.new("RGB", (100, 100))), 50, 400, 200, 200)
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def test_txt_extraction_and_passages():
    data = (FIXTURES / "breve.txt").read_bytes()
    doc = extract_file("breve.txt", data, "trascrizione")
    build_transcript_passages(doc)
    assert "Orizzonte Famiglia" in doc.text
    assert doc.passages[0].id == "T001"
    assert "".join(p.text for p in doc.passages).count("25.000") == 1


def test_txt_cp1252_encoding():
    data = "Il caricamento è dell'1,5% e la penale è del 2%.".encode("cp1252")
    doc = extract_file("vecchio.txt", data, "trascrizione")
    assert "è" in doc.text


def test_docx_extraction():
    import docx

    document = docx.Document()
    document.add_paragraph("Relatore: il premio minimo è di 1.200 euro l'anno.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Profilo 1"
    table.rows[0].cells[1].text = "80% fondi"
    buffer = io.BytesIO()
    document.save(buffer)
    doc = extract_file("lezione.docx", buffer.getvalue(), "trascrizione")
    assert "1.200 euro" in doc.text and "Profilo 1 | 80% fondi" in doc.text


def test_text_pdf_extraction():
    data = _text_pdf(["Relatore: la durata del piano e' di dieci anni.", "Il premio minimo e' 100 euro al mese."])
    doc = extract_file("lezione.pdf", data, "trascrizione")
    assert "dieci anni" in doc.text
    assert doc.pages_without_text == []


def test_image_only_pdf_is_rejected_clearly():
    with pytest.raises(ExtractionError) as info:
        extract_file("scansione.pdf", _image_only_pdf(), "trascrizione")
    assert "soltanto immagini" in str(info.value)


def test_pdf_with_some_image_pages_warns():
    data = _text_pdf(["Relatore: testo della prima pagina con contenuto sufficiente."], extra_image_page=True)
    doc = extract_file("misto.pdf", data, "trascrizione")
    assert doc.pages_without_text == [2]
    assert any("pagine 2" in w for w in doc.warnings)


def test_unsupported_and_empty():
    with pytest.raises(ExtractionError):
        extract_file("file.doc", b"xxx", "trascrizione")
    with pytest.raises(ExtractionError):
        from_pasted_text("   ")


def test_truncation_detected():
    text = (FIXTURES / "interrotta.txt").read_text(encoding="utf-8")
    signals = detect_issues(text)
    assert any("limite di durata" in s for s in signals)
    assert any("invito a un abbonamento" in s for s in signals)


def test_unterminated_ending_detected():
    signals = detect_issues("Il relatore spiega la copertura e poi dice che il limite per i danni da acqua è pari a")
    assert any("senza punteggiatura" in s for s in signals)


def test_complete_text_has_no_truncation_signal():
    text = (FIXTURES / "breve.txt").read_text(encoding="utf-8")
    assert not any("interruzione" in s.lower() for s in detect_issues(text))


def test_long_segmentation_keeps_final_part():
    text = (FIXTURES / "lunga.txt").read_text(encoding="utf-8")
    doc = from_pasted_text(text)
    build_transcript_passages(doc)
    segments = segment_passages(doc.passages, 40000)
    assert len(segments) >= 4
    covered = {p.id for segment in segments for p in segment}
    assert covered == {p.id for p in doc.passages}
    assert "costa 50 euro" in segments[-1][-1].text or "costa 50 euro" in " ".join(p.text for p in segments[-1])
    # Sovrapposizione di un passaggio tra segmenti consecutivi.
    for previous, current in zip(segments, segments[1:]):
        assert previous[-1].id == current[0].id
