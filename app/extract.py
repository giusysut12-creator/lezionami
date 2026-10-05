"""Fase 1: estrazione e controllo del testo.

- Legge TXT, DOCX e PDF testuali.
- Riconosce i PDF composti solo da immagini (scansioni) invece di elaborare testo vuoto.
- Rileva segnali di trascrizione interrotta o incompleta.
- Divide il testo in passaggi numerati (T001, T002...) usati come riferimenti interni.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import PurePath


class ExtractionError(Exception):
    """Errore comprensibile per l'utente durante l'estrazione."""


@dataclass
class Passage:
    id: str
    text: str
    page: int | None = None


@dataclass
class ExtractedDocument:
    name: str
    kind: str  # "trascrizione" | "fonte"
    file_type: str
    text: str
    pages_total: int | None = None
    pages_without_text: list[int] = field(default_factory=list)
    metadata_title: str = ""
    metadata_date: str = ""
    warnings: list[str] = field(default_factory=list)
    passages: list[Passage] = field(default_factory=list)
    page_texts: list[str] = field(default_factory=list)


TEXT_EXTENSIONS = {".txt", ".text", ".md", ".srt", ".vtt"}
SUPPORTED = TEXT_EXTENSIONS | {".docx", ".pdf"}
MIN_PAGE_CHARS = 25


def _decode_text(data: bytes) -> str:
    encodings = ["utf-8-sig", "cp1252", "latin-1"]
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        encodings.insert(0, "utf-16")
    for encoding in encodings:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Impossibile leggere la codifica del file di testo.")


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = text.replace("﻿", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _clean_pdf_page(text: str) -> str:
    # Rimuove sequenze "(cid:123)" prodotte da font non mappati.
    text = re.sub(r"\(cid:\d+\)", "", text)
    # Riunisce parole spezzate a fine riga con trattino.
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    return text


def _extract_pdf(data: bytes, doc: ExtractedDocument) -> None:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError("Libreria pypdf non installata.") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ExtractionError(
                    f"Il PDF «{doc.name}» è protetto da password e non può essere letto."
                ) from exc
        pages = reader.pages
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"Il file «{doc.name}» non è un PDF leggibile.") from exc

    doc.pages_total = len(pages)
    try:
        meta = reader.metadata or {}
        doc.metadata_title = str(meta.get("/Title") or "").strip()
        raw_date = str(meta.get("/CreationDate") or "")
        match = re.search(r"D:(\d{4})(\d{2})(\d{2})", raw_date)
        if match:
            doc.metadata_date = f"{match.group(3)}/{match.group(2)}/{match.group(1)}"
    except Exception:
        pass

    texts: list[str] = []
    for number, page in enumerate(pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        page_text = _normalize(_clean_pdf_page(page_text))
        if len(re.sub(r"\s", "", page_text)) < MIN_PAGE_CHARS:
            doc.pages_without_text.append(number)
        texts.append(page_text)
    doc.page_texts = texts
    doc.text = "\n\n".join(t for t in texts if t)

    if doc.pages_total == 0:
        raise ExtractionError(f"Il PDF «{doc.name}» non contiene pagine.")
    if len(doc.pages_without_text) == doc.pages_total:
        raise ExtractionError(
            f"Il PDF «{doc.name}» contiene soltanto immagini (per esempio una scansione): "
            "non c'è testo selezionabile da elaborare. Carica un PDF testuale, un file TXT/DOCX, "
            "oppure esegui prima il riconoscimento del testo (OCR)."
        )
    if doc.pages_without_text:
        pages_list = ", ".join(str(p) for p in doc.pages_without_text[:20])
        more = "…" if len(doc.pages_without_text) > 20 else ""
        doc.warnings.append(
            f"Nel PDF «{doc.name}» le pagine {pages_list}{more} non contengono testo "
            "selezionabile (probabilmente immagini): il loro contenuto non viene elaborato."
        )


def _extract_docx(data: bytes, doc: ExtractedDocument) -> None:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError("Libreria python-docx non installata.") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(
            f"Il file «{doc.name}» non è un DOCX valido (i vecchi file .doc non sono supportati)."
        ) from exc
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    try:
        core = document.core_properties
        doc.metadata_title = (core.title or "").strip()
        if core.created:
            doc.metadata_date = core.created.strftime("%d/%m/%Y")
    except Exception:
        pass
    doc.text = _normalize("\n".join(parts))


def extract_file(name: str, data: bytes, kind: str) -> ExtractedDocument:
    suffix = PurePath(name or "").suffix.lower()
    if suffix not in SUPPORTED:
        raise ExtractionError(
            f"Formato non supportato per «{name}». Usa un file TXT, DOCX o PDF testuale."
        )
    doc = ExtractedDocument(name=name, kind=kind, file_type=suffix.lstrip("."), text="")
    if suffix == ".pdf":
        _extract_pdf(data, doc)
    elif suffix == ".docx":
        _extract_docx(data, doc)
    else:
        doc.text = _normalize(_decode_text(data))
    if len(doc.text.strip()) < 20:
        raise ExtractionError(f"Il file «{name}» non contiene testo sufficiente da elaborare.")
    return doc


def from_pasted_text(text: str) -> ExtractedDocument:
    clean = _normalize(text or "")
    if len(clean) < 20:
        raise ExtractionError("Il testo incollato è vuoto o troppo breve.")
    return ExtractedDocument(name="testo incollato", kind="trascrizione", file_type="testo", text=clean)


# ---------------------------------------------------------------------------
# Rilevamento di interruzioni e limiti della trascrizione
# ---------------------------------------------------------------------------

TRUNCATION_PATTERNS = [
    (r"(questo|il)\s+file\s+(è|e'|e)\s+più\s+lung[oa]\s+di\s+\d+\s+minuti", "messaggio del servizio di trascrizione su un limite di durata"),
    (r"this\s+file\s+is\s+longer\s+than\s+\d+\s+minutes", "messaggio del servizio di trascrizione su un limite di durata"),
    (r"(solo|soltanto)\s+i\s+primi\s+\d+\s+minuti", "indicazione che è stata trascritta solo una parte"),
    (r"only\s+the\s+first\s+\d+\s+minutes", "indicazione che è stata trascritta solo una parte"),
    (r"upgrade\s+to\s+(unlimited|pro|premium)", "invito a un abbonamento del servizio di trascrizione"),
    (r"passa\s+(a|al)\s+(piano\s+)?(unlimited|pro|premium)", "invito a un abbonamento del servizio di trascrizione"),
    (r"\bturboscribe\b", "riferimento al servizio di trascrizione TurboScribe"),
    (r"trascrizione\s+(troncata|interrotta|parziale|incompleta)", "dichiarazione di trascrizione incompleta"),
    (r"transcript(ion)?\s+(truncated|cut\s+off|incomplete)", "dichiarazione di trascrizione incompleta"),
]

UNCLEAR_PATTERNS = [
    r"\[(inaudible|incomprensibile|inudibile|non\s+udibile|\?\?\?|unclear)\]",
    r"\(inaudible\)|\(incomprensibile\)",
]


def detect_issues(text: str) -> list[str]:
    signals: list[str] = []
    lowered = text.lower()
    seen: set[str] = set()
    for pattern, label in TRUNCATION_PATTERNS:
        match = re.search(pattern, lowered)
        if match and label not in seen:
            seen.add(label)
            position = match.start() / max(len(lowered), 1)
            where = "alla fine" if position > 0.8 else ("all'inizio" if position < 0.2 else "nel corpo")
            signals.append(
                f"Possibile interruzione: {label} («{text[match.start():match.end()]}», {where} del testo)."
            )
    unclear = sum(len(re.findall(p, lowered)) for p in UNCLEAR_PATTERNS)
    if unclear:
        signals.append(f"La trascrizione contiene {unclear} passaggi segnati come incomprensibili.")
    tail = text.rstrip()
    if tail:
        last_line = tail.splitlines()[-1].strip()
        if last_line and last_line[-1] not in ".?!…»\"')]" and len(last_line) > 25:
            signals.append(
                "Il testo termina senza punteggiatura finale: l'ultima frase potrebbe essere interrotta."
            )
    return signals


# ---------------------------------------------------------------------------
# Suddivisione in passaggi con identificativi stabili
# ---------------------------------------------------------------------------

_SENTENCE_END = re.compile(r"(?<=[.!?…])\s+")


def _split_long(text: str, max_chars: int) -> list[str]:
    if len(text) <= max_chars:
        return [text]
    sentences = _SENTENCE_END.split(text)
    chunks: list[str] = []
    current = ""
    for sentence in sentences:
        while len(sentence) > max_chars:  # frase lunghissima senza punteggiatura
            cut = sentence.rfind(" ", 0, max_chars)
            cut = cut if cut > max_chars // 2 else max_chars
            if current:
                chunks.append(current)
                current = ""
            chunks.append(sentence[:cut].strip())
            sentence = sentence[cut:].strip()
        if current and len(current) + len(sentence) + 1 > max_chars:
            chunks.append(current)
            current = sentence
        else:
            current = f"{current} {sentence}".strip()
    if current:
        chunks.append(current)
    return chunks


def make_passages(text: str, prefix: str, target_chars: int = 900, max_chars: int = 1400,
                  page: int | None = None, start_index: int = 1) -> list[Passage]:
    blocks = [b.strip() for b in re.split(r"\n\s*\n|\n", text) if b.strip()]
    pieces: list[str] = []
    current = ""
    for block in blocks:
        for piece in _split_long(block, max_chars):
            if current and len(current) + len(piece) + 1 > target_chars:
                pieces.append(current)
                current = piece
            else:
                current = f"{current}\n{piece}".strip()
    if current:
        pieces.append(current)
    return [
        Passage(id=f"{prefix}{index:03d}", text=piece, page=page)
        for index, piece in enumerate(pieces, start=start_index)
    ]


def build_transcript_passages(doc: ExtractedDocument) -> None:
    doc.passages = make_passages(doc.text, "T")


def build_source_passages(doc: ExtractedDocument, source_index: int) -> None:
    prefix = f"D{source_index}-"
    if doc.page_texts:
        passages: list[Passage] = []
        counter = 1
        for page_number, page_text in enumerate(doc.page_texts, start=1):
            if not page_text.strip():
                continue
            page_passages = make_passages(page_text, prefix, page=page_number, start_index=counter)
            counter += len(page_passages)
            passages.extend(page_passages)
        doc.passages = passages
    else:
        doc.passages = make_passages(doc.text, prefix)


def render_passages(passages: list[Passage]) -> str:
    lines = []
    for passage in passages:
        page = f" p.{passage.page}" if passage.page else ""
        lines.append(f"[{passage.id}{page}] {passage.text}")
    return "\n\n".join(lines)


def segment_passages(passages: list[Passage], max_chars: int) -> list[list[Passage]]:
    """Raggruppa i passaggi in segmenti per l'inventario delle trascrizioni lunghe.

    Ogni segmento ripete l'ultimo passaggio del precedente per non perdere il
    contesto di frasi che attraversano il confine.
    """
    segments: list[list[Passage]] = []
    current: list[Passage] = []
    size = 0
    for passage in passages:
        if current and size + len(passage.text) > max_chars:
            segments.append(current)
            current = [current[-1]]
            size = len(current[0].text)
        current.append(passage)
        size += len(passage.text)
    if current:
        segments.append(current)
    return segments
